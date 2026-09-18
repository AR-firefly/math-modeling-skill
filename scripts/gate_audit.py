#!/usr/bin/env python3
"""Read-only project artifact audit. Use --project DIR, or isolated --self-test.
Machine checks never replace independent semantic review or award judgment.
"""
import argparse
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))


def _project_file(project, name):
    if not isinstance(name, str) or not name or Path(name).is_absolute():
        raise ValueError("Evidence path must be project-relative")
    path = (project / name).resolve()
    if project not in path.parents or not path.is_file():
        raise ValueError(f"Missing or outside-project evidence: {name}")
    return path


def validate_figure_source(figure):
    """Check recorded publisher evidence; independent review verifies its truth."""
    from urllib.parse import urlparse
    issues = []
    url = figure.get("official_url", "")
    evidence = figure.get("source_evidence")
    excluded = ("csdn.net", "baidu.com", "xiaohongshu.com", "zhihu.com",
                "weibo.com", "toutiao.com", "bilibili.com", "jianshu.com", "data-to-viz.com")
    for candidate in (url, evidence.get("original_url", "") if isinstance(evidence, dict) else ""):
        parsed = urlparse(candidate)
        host = (parsed.hostname or "").lower()
        if parsed.scheme not in ("https", "http") or not host:
            issues.append("Figure source must identify an original official HTTP(S) document")
        if any(host == domain or host.endswith("." + domain) for domain in excluded):
            issues.append("Excluded figure documentation source: " + host)
        if host in ("github.com", "raw.githubusercontent.com") and "/prompt-library" in parsed.path:
            issues.append("TRIZ exception is not plotting API documentation")
    fields = ("publisher", "original_url", "purpose", "verification_note")
    if (not isinstance(evidence, dict) or evidence.get("status") != "verified"
            or evidence.get("original_source") is not True
            or any(not isinstance(evidence.get(k), str) or not evidence[k].strip() for k in fields)):
        issues.append("Figure needs verified publisher/original-source evidence, not a domain alone")
    elif evidence["original_url"] != url:
        issues.append("Figure official URL does not match publisher evidence")
    return issues


def validate_review(project, review, artifact_names, *, frozen_criteria_hash=None):
    import hashlib
    project = Path(project).resolve()
    issues = []
    if not isinstance(review, dict) or review.get("schema_version") != 1:
        return ["Final review schema_version must be 1"]
    try:
        name = review.get("criteria_file")
        if name != "references/终审运行检查表.md":
            issues.append("Final review criteria_file must identify frozen criteria")
        criteria = _project_file(project, name)
        if frozen_criteria_hash is not None and hashlib.sha256(criteria.read_bytes()).hexdigest() != frozen_criteria_hash:
            issues.append("Final review differs from frozen criteria baseline")
        if hashlib.sha256(criteria.read_bytes()).hexdigest() != review.get("criteria_sha256"):
            issues.append("Final review criteria hash mismatch")
    except (ValueError, OSError) as exc:
        issues.append(str(exc))
    reviewers = []
    review_evidence = set()
    for role in ("self_review", "independent_review"):
        part = review.get(role)
        if not isinstance(part, dict):
            issues.append(f"Missing {role}")
            continue
        reviewer = part.get("reviewer")
        if not isinstance(reviewer, str) or not reviewer.strip():
            issues.append(f"{role}: missing reviewer")
        reviewers.append(reviewer)
        if part.get("decision") != "passed" or part.get("open_issues") != []:
            issues.append(f"{role}: unresolved review")
        evidence = part.get("evidence")
        if not isinstance(evidence, list) or not evidence:
            issues.append(f"{role}: missing evidence")
        else:
            for name in evidence:
                try:
                    _project_file(project, name)
                    review_evidence.add(name)
                except (ValueError, OSError) as exc:
                    issues.append(f"{role}: {exc}")
    if len(reviewers) == 2 and reviewers[0] == reviewers[1]:
        issues.append("Independent reviewer must differ from self reviewer")
    hashes = review.get("artifact_hashes", {})
    if not isinstance(hashes, dict):
        return issues + ["artifact_hashes must be a mapping"]
    for name in set(artifact_names) | set(hashes) | review_evidence:
        try:
            path = _project_file(project, name)
            if hashlib.sha256(path.read_bytes()).hexdigest() != hashes.get(name):
                issues.append(f"Missing or stale review artifact hash: {name}")
        except (ValueError, OSError) as exc:
            issues.append(str(exc))
    return issues


def tex_readable_body(tex):
    """Conservative extraction for this renderer; unknown content stays visible.

    Remove preamble and known layout arguments, never arbitrary numeric commands
    or mathematical expressions. This is not a general TeX interpreter.
    """
    text = tex.split(r"\begin{document}", 1)[-1]
    text = text.split(r"\section{人工智能使用声明}", 1)[0]
    text = re.sub(r"(?<!\\)%[^\n]*", "", text)
    text = re.sub(r"\\includegraphics\*?(?:\[[^\]]*\])?\{[^}]*\}", "", text)
    text = re.sub(r"\\(?:vspace|hspace)\*?\{[^}]*\}", "", text)
    text = re.sub(r"\\\\(?:\[[^\]]*\])?", " ", text)
    text = re.sub(r"\\(?:begin|end)\{(?:table|figure|center|tabular)\}(?:\[[^\]]*\])?", " ", text)
    text = re.sub(r"\\pagenumbering\{[^}]*\}", "", text)
    text = re.sub(r"\\([_%&#{}$])", r"\1", text)
    # Decode named literal escapes emitted by tex_renderer after layout parsing.
    # Inserting a literal backslash earlier could accidentally form a TeX command.
    for escaped, literal in ((r"\textasciitilde{}", "~"), (r"\textasciicircum{}", "^"),
                             (r"\textbackslash{}", "\\")):
        text = text.replace(escaped, literal)
    return text


def verify_artifact_numbers(project, paper, doc_body, tex):
    """Verify both mapped and unclassified numbers in structured and actual text."""
    from math_modeling.verify import verify_paper_numbers
    project = Path(project)
    def flatten(node):
        if isinstance(node, dict):
            return "\n".join(flatten(value) for key, value in node.items()
                             if key not in ("references", "meta", "path") and not key.startswith("_"))
        if isinstance(node, (list, tuple)):
            return "\n".join(flatten(value) for value in node)
        return str(node)
    claims = json.loads((project / "output/claims.json").read_text(encoding="utf-8"))
    exception_path = project / "output/non_result_numbers.json"
    exceptions = json.loads(exception_path.read_text(encoding="utf-8")) if exception_path.exists() else []
    if not isinstance(exceptions, list):
        return ["non_result_numbers.json must contain a list"]
    issues = []
    bodies = (("JSON", flatten(paper)), ("DOCX", doc_body), ("TeX", tex_readable_body(tex)))
    for label, body in bodies:
        # A classification can concern a title or other text appearing only in a
        # rendered artifact. Validate it in the applicable representation.
        relevant = [entry for entry in exceptions if not isinstance(entry, dict)
                    or not isinstance(entry.get("text"), str) or entry["text"] in body]
        found, unmatched = verify_paper_numbers(project / "results/results.json", body,
            claims=claims, reverse=True, non_result_numbers=relevant)
        issues.extend(label + ": " + item for item in found)
        if unmatched:
            issues.append(f"{label}: unclassified numeric claims: {', '.join(unmatched)}")
    for entry in exceptions:
        if isinstance(entry, dict) and isinstance(entry.get("text"), str) and not any(entry["text"] in body for _, body in bodies):
            issues.append("Non-result numeric classification not found in any artifact: " + entry["text"])
    return issues


def audit_project(project):
    project = Path(project).resolve()
    issues = []
    report = {"machine_checks_only": True, "independent_semantic_review_required": True,
              "project": str(project), "passed": False, "issues": issues}
    if project == ROOT or ROOT in project.parents:
        issues.append("SKILL_ROOT cannot be audited as a contest project")
        return report
    required = ["output/paper.json", "output/paper.tex", "output/paper.docx",
                "output/claims.json", "output/source_evidence.json", "output/process_record.md",
                "output/final_review.json", "results/results.json", "figures/figures_manifest.json",
                "output/workflow_state.json"]
    for name in required:
        if not (project / name).is_file():
            issues.append("Missing: " + name)
    if issues:
        return report
    try:
        def read(name):
            return json.loads((project / name).read_text(encoding="utf-8"))
        from math_modeling.workflow import validate_current_approval
        state = validate_current_approval(project)
        from math_modeling.workflow import validate_stage2_outputs
        validate_stage2_outputs(project, state)
        if state.get("status") not in ("final_awaiting_review", "ai_complete"):
            issues.append("Real project has not completed stage2 for final review")
        if not state.get("criteria_sha256"):
            issues.append("Missing frozen criteria baseline in real workflow state")
        paper = read("output/paper.json")
        claims = read("output/claims.json")
        sources = read("output/source_evidence.json")
        figures = read("figures/figures_manifest.json")
        review = read("output/final_review.json")
        if paper.get("_placeholder_issues"):
            issues.append("paper.json contains unresolved placeholders")
        if paper.get("ai_declaration", ""):
            issues.append("AI declaration must remain blank for the team")
        from math_modeling.paper_generator import find_paper_placeholders
        if find_paper_placeholders(paper):
            issues.append("Paper contains unresolved placeholder text")
        from math_modeling.paper_generator import validate_references
        from math_modeling.verify import verify_paper_numbers, verify_citations, verify_figure_references
        if not isinstance(sources, list):
            issues.append("source_evidence.json must contain a list")
            sources = []
        issues.extend(validate_references(paper.get("references", []), source_evidence=sources))
        tex = (project / "output/paper.tex").read_text(encoding="utf-8")
        from docx import Document
        doc = Document(project / "output/paper.docx")
        paragraphs = [p.text for p in doc.paragraphs]
        if "人工智能使用声明" not in paragraphs or "参考文献" not in paragraphs:
            issues.append("DOCX missing human declaration or references heading")
        else:
            start, end = paragraphs.index("人工智能使用声明"), paragraphs.index("参考文献")
            if start >= end or any(t.strip() for t in paragraphs[start+1:end]):
                issues.append("DOCX declaration must be blank before references")
        declaration = re.search(r"\\section\{人工智能使用声明\}(.*?)\\section\{参考文献\}", tex, re.S)
        if not declaration or re.sub(r"\\vspace\{[^}]*\}|\s+", "", declaration.group(1)):
            issues.append("TeX declaration must be blank before references")
        # Inspect actual rendered artifacts, never synthesize prose from results.
        doc_body = "\n".join(paragraphs[:paragraphs.index("人工智能使用声明")] if "人工智能使用声明" in paragraphs else paragraphs)
        doc_body += "\n" + "\n".join(" ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
        tex_body = tex.split(r"\section{人工智能使用声明}")[0]
        tex_body = re.sub(r"\\([_%&#{}])", r"\1", tex_body)
        if not isinstance(claims, list) or not claims:
            issues.append("Missing nonempty result claim mapping")
        else:
            issues.extend(verify_artifact_numbers(project, paper, doc_body, tex))
            issues.extend(verify_citations(doc_body, paper.get("references", [])))
        if not isinstance(figures, list):
            raise ValueError("figure manifest must be a list")
        if figures or any(section.get("images") for section in paper.get("sections", [])):
            issues.extend(verify_figure_references(paper, figures_dir=str(project / "figures")))
        hashed_artifacts = [name for name in required if name != "output/final_review.json"]
        if (project / "output/non_result_numbers.json").is_file():
            hashed_artifacts.append("output/non_result_numbers.json")
        for figure in figures:
            no = figure.get("no", "unknown")
            issues.extend(f"{no}: {item}" for item in validate_figure_source(figure))
            for key in ("script", "data_entrypoint", "run_command", "dependencies", "official_url",
                        "official_api", "adaptation", "changes", "reason", "title", "conclusion"):
                if not figure.get(key):
                    issues.append(f"{no}: missing {key}")
            for key, prefix in (("file", "figures"), ("script", ""), ("data_entrypoint", "")):
                value = figure.get(key)
                try:
                    relative = (Path(prefix) / value).as_posix() if isinstance(value, str) else ""
                    _project_file(project, relative)
                    hashed_artifacts.append(relative)
                except (ValueError, OSError):
                    issues.append(f"{no}: missing local {key}")
            if figure.get("evidence_status") != "verified":
                issues.append(f"{no}: reproduction and source evidence not verified")
        record = (project / "output/process_record.md").read_text(encoding="utf-8")
        if len(record.splitlines()) < 500 or len(re.findall(r"^## ", record, re.M)) < 13:
            issues.append("Process record does not meet existing 500-line / 13-section contract")
        issues.extend(validate_review(project, review, hashed_artifacts,
                                      frozen_criteria_hash=state.get("criteria_sha256")))
    except Exception as exc:
        issues.append(f"Invalid artifact: {type(exc).__name__}: {exc}")
    report["passed"] = not issues
    return report


def self_test():
    """Copy only runnable assets; all writes and child cwd stay in a temporary tree."""
    with tempfile.TemporaryDirectory(prefix="math_skill_selftest_") as temp:
        work = Path(temp) / "skill"
        work.mkdir()
        demo_project = Path(temp) / "demo-project"
        demo_project.mkdir()
        for folder in ("src", "tests", "examples", "scripts", "algorithms"):
            source = ROOT / folder
            if source.exists():
                shutil.copytree(source, work / folder, ignore=shutil.ignore_patterns("__pycache__", "results", "figures", "output", ".pytest_cache"))
        (work / "references").mkdir(exist_ok=True)
        for reference in (ROOT / "references").glob("*.md"):
            shutil.copy2(reference, work / "references" / reference.name)
        shutil.copy2(ROOT / "run_all.py", work / "run_all.py")
        env = dict(os.environ, PYTHONPATH=os.pathsep.join((str(work), str(work / "src"))), PYTEST_DISABLE_PLUGIN_AUTOLOAD="1", PYTHONDONTWRITEBYTECODE="1", MPLCONFIGDIR=str(work / "mpl"))
        commands = [[sys.executable, "-B", "-m", "pytest", "tests", "-q", "-p", "no:cacheprovider"],
                    [sys.executable, "-B", str(work / "run_all.py"), "--demo"],
                    [sys.executable, "-B", "scripts/validate_figures.py", "--self-test"]]
        commands += [[sys.executable, "-B", str(work / "examples" / script)] for script in
                     ("main.py", "评价_TOPSIS_demo.py", "预测_GM11_demo.py", "机理_ODE_demo.py")]
        results = []
        for command in commands:
            cwd = demo_project if "--demo" in command else (work / "examples" if "examples" in str(command[-1]) else work)
            result = subprocess.run(command, cwd=cwd, env=env, capture_output=True, text=True,
                                    encoding="utf-8", errors="replace", timeout=600)
            results.append({"command": command, "exit_code": result.returncode,
                            "output": (result.stdout + result.stderr)[-2000:]})
        return {"self_test_only": True, "project_verified": False,
                "passed": all(r["exit_code"] == 0 for r in results), "checks": results}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--project", type=Path)
    group.add_argument("--self-test", action="store_true")
    args = parser.parse_args(argv)
    result = self_test() if args.self_test else audit_project(args.project)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if result["passed"] else 1


if __name__ == "__main__":
    sys.exit(main())
