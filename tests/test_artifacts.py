import json
import urllib.error
import pytest
from math_modeling.paper_generator import build_paper_content, validate_references
from math_modeling.visualizer import Visualizer
from math_modeling.tex_renderer import render_tex
from math_modeling.docx_renderer import render_docx
from docx import Document


def test_real_references_and_placeholders(tmp_path):
    p = tmp_path / "results.json"
    p.write_text(json.dumps({"Q1_cost": 25}))
    refs = ["Author. Paper. 2024. DOI:10.1000/example"]
    paper = build_paper_content(p, {"title": "Example", "references": refs}, ref_dir=None)
    assert paper["references"] == refs
    assert paper["_placeholder_issues"]
    assert paper["ai_declaration"] == ""


def test_render_human_blank_before_references(tmp_path):
    paper = {"meta": {"title": "Title"}, "references": ["Reference"], "ai_declaration": ""}
    tex = render_tex(paper)
    assert tex.index("人工智能使用声明") < tex.index("参考文献")
    path = tmp_path / "paper.docx"
    render_docx(paper, path)
    paras = [p.text for p in Document(path).paragraphs]
    i = paras.index("人工智能使用声明")
    assert paras[i + 1] == ""
    assert i < paras.index("参考文献")


def test_old_figure_is_not_evidence_complete(tmp_path):
    v = Visualizer(output_dir=tmp_path)
    v.register_figure("fig1", "title", "claim", "results.json:Q1")
    assert v._figures[0]["evidence_status"] == "incomplete"


def test_figure_evidence_fields(tmp_path):
    v = Visualizer(output_dir=tmp_path)
    v.register_figure("fig1", "title", "claim", "results.json:Q1",
        script="plot.py", data_entrypoint="data.csv", run_command="python plot.py",
        dependencies={"matplotlib": "3.10"}, official_url="https://matplotlib.org/stable/api/_as_gen/matplotlib.pyplot.plot.html",
        official_api="matplotlib.pyplot.plot", adaptation="official_api_composition",
        changes="Use measured values", reason="Show trend", seed=None)
    record=v._figures[0]
    assert record["script"] == "plot.py"
    assert record["evidence_status"] == "pending_verification"


def test_offline_doi_not_verified(monkeypatch):
    import urllib.request
    def offline(*args, **kwargs):
        raise urllib.error.URLError("offline")
    monkeypatch.setattr(urllib.request, "urlopen", offline)
    assert validate_references(["A. 2024. DOI:10.1000/example"], check_doi=True)


def test_excluded_reference():
    assert validate_references(["A. 2024. https://blog.csdn.net/a/article/details/123"])


def test_legacy_declaration_empty():
    from math_modeling.paper_generator import PaperGenerator, PaperConfig
    gen = PaperGenerator(PaperConfig())
    gen.write_ai_report()
    assert [p.text for p in gen.doc.paragraphs] == ["人工智能使用声明", ""]


def test_source_evidence_cannot_be_missing():
    assert validate_references(["A. 2024. https://example.org/paper"], source_evidence=[])


def test_supplied_complete_sections_and_keywords(tmp_path):
    p=tmp_path / "results.json"
    p.write_text("{}")
    sections=[{"title":"Model", "paras":["Evidence"], "images":[], "tables":[], "formulas":[]}]
    paper=build_paper_content(p, {"title":"Title", "abstract":["Constructed a flow model."],
        "sections":sections, "keywords":["network flow"], "references":["A. 2024. https://example.org"],
        "appendix":["solve.py"]}, ref_dir=None)
    assert paper["sections"] == sections
    assert "network flow" in render_tex(paper)
    assert not paper["_placeholder_issues"]


def test_figure_register_preserves_publisher_evidence(tmp_path):
    evidence={"publisher":"Seaborn maintainers","original_source":True,"status":"verified",
              "original_url":"https://seaborn.pydata.org/generated/seaborn.lineplot.html", "purpose":"plotting", "verification_note":"Checked official docs"}
    v=Visualizer(output_dir=tmp_path)
    v.register_figure("fig1","Trend","Rising","data.csv",source_evidence=evidence)
    assert v._figures[0]["source_evidence"] == evidence
    assert v._figures[0]["evidence_status"] == "incomplete"


def test_math_comparisons_are_not_placeholders(tmp_path):
    p=tmp_path/"results.json"; p.write_text("{}",encoding="utf-8")
    meta={"title":"模型", "abstract":["约束 x<1 且 y>2"], "sections":[], "references":["A. 2024. https://example.org"], "appendix":["源程序已交付"]}
    paper=build_paper_content(p,meta,ref_dir=None)
    assert not paper["_placeholder_issues"]


def test_ai_instruction_placeholder_detected():
    from math_modeling.paper_generator import find_paper_placeholders
    assert find_paper_placeholders({"sections":[{"paras":["<AI 依据赛题生成>"]}]})
    assert not find_paper_placeholders({"sections":[{"paras":["0 < x < 1"]}]})
