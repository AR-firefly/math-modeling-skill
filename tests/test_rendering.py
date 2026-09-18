from docx import Document
from docx.oxml.ns import qn
from math_modeling.tex_renderer import render_tex, _render_image
from math_modeling.docx_renderer import render_docx


def test_abstract_page_break_before_numbering():
    tex=render_tex({"meta":{"title":"Title"}})
    assert tex.index(r"\newpage") < tex.index(r"\pagenumbering{arabic}")


def test_notes_and_three_line_table(tmp_path):
    paper={"meta":{"title":"Title"},"sections":[{"title":"Results", "paras":["The table supports the result."],
        "tables":[{"caption":"Table 1", "headers":["Metric","Value"], "rows":[["Cost",25]],"notes":"Values in yuan."}]}]}
    tex=render_tex(paper)
    assert "Values in yuan." in tex
    assert r"\toprule" in tex and r"\midrule" in tex and r"\bottomrule" in tex
    path=tmp_path/"paper.docx"
    render_docx(paper,path)
    doc=Document(path)
    assert "Values in yuan." in [p.text for p in doc.paragraphs]
    table=doc.tables[0]
    assert table.style.name != "Table Grid"
    borders=table._tbl.tblPr.find(qn("w:tblBorders"))
    assert borders.find(qn("w:insideV")).get(qn("w:val")) == "nil"
    assert borders.find(qn("w:top")).get(qn("w:val")) == "single"
    assert table.rows[0].cells[0]._tc.tcPr.find(qn("w:tcBorders")).find(qn("w:bottom")).get(qn("w:val")) == "single"


def test_image_notes_tex():
    assert "Measured at noon" in _render_image({"path":"image.png","caption":"Figure 1", "notes":["Measured at noon"]})


def test_docx_heading_color_and_body_page_number(tmp_path):
    path=tmp_path/"paper.docx"
    render_docx({"meta":{"title":"Title"},"abstract":["Abstract"],"sections":[]},path)
    doc=Document(path)
    assert str(doc.styles["Heading 1"].font.color.rgb) == "000000"
    assert len(doc.sections) == 2
    assert "PAGE" not in doc.sections[0].footer._element.xml
    assert "PAGE" in doc.sections[1].footer._element.xml
    assert doc.sections[1]._sectPr.find(qn("w:pgNumType")).get(qn("w:start")) == "1"


def test_comprehensive_render_fixture(tmp_path, monkeypatch):
    import matplotlib.pyplot as plt
    import shutil
    import tempfile
    from pathlib import Path
    monkeypatch.chdir(tmp_path)
    fig, ax=plt.subplots()
    ax.plot([0,1,2],[0,1,4])
    ax.set_xlabel("x")
    ax.set_ylabel("y")
    image=tmp_path/"trend.png"
    fig.savefig(image,dpi=150)
    plt.close(fig)
    paper={"meta":{"title":"二次增长模型验证", "keywords":["二次模型","误差分析"]},
        "abstract":["建立二次模型解释增长趋势，并采用误差分析检验数值结果。"],
        "sections":[{"title":"模型与结果", "formulas":[r"y=x^2"],
        "images":[{"path":str(image),"caption":"图1 二次增长趋势", "notes":"横轴为时间，纵轴为无量纲响应。"}],
        "tables":[{"caption":"表1 模型预测结果","headers":["时间","响应"],"rows":[[0,0],[1,1],[2,4]],"notes":"响应为无量纲变量。"}],
        "paras":["图1与表1显示，时间增加时响应按平方规律增长，模型预测与计算结果一致。"]}],
        "references":["参考来源由实际赛题核验后填写，此处仅为排版回归样例。"]}
    out=tmp_path/"comprehensive.docx"
    render_docx(paper,out)
    doc=Document(out)
    assert len(doc.inline_shapes) == 2
    stable=tmp_path/"qa_copy"
    stable.mkdir(exist_ok=True)
    shutil.copy2(out,stable/out.name)
    assert "横轴为时间，纵轴为无量纲响应。" in [p.text for p in doc.paragraphs]


def test_table_notes_stay_with_last_row(tmp_path):
    out=tmp_path/"table.docx"
    render_docx({"meta":{"title":"Title"},"sections":[{"title":"Results","tables":[{"headers":["x"],"rows":[[1]],"caption":"Table","notes":["Units","Condition"]}]}]},out)
    doc=Document(out)
    assert doc.tables[0].rows[-1].cells[0].paragraphs[0].paragraph_format.keep_with_next
    notes=[p for p in doc.paragraphs if p.text == "Units"]
    assert notes[0].paragraph_format.keep_with_next


def test_short_tables_stay_together_long_tables_repeat_header(tmp_path):
    for count in (3,20):
        out=tmp_path/f"table{count}.docx"
        render_docx({"meta":{"title":"Title"},"sections":[{"title":"Results","tables":[{"headers":["x"],"rows":[[i] for i in range(count)],"caption":"Table","notes":"Units"}]}]},out)
        table=Document(out).tables[0]
        assert table.rows[0]._tr.trPr.find(qn("w:tblHeader")) is not None
        keep=table.rows[1].cells[0].paragraphs[0].paragraph_format.keep_with_next
        assert bool(keep) == (count <= 10)
